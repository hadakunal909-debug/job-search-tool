"""
export.py — turn plain-text résumés / cover letters into a downloadable .docx.

Light, dependency-isolated: if python-docx isn't installed, build_docx raises ImportError and
the caller falls back to serving plain text, so the app never hard-crashes on a missing dep.

Heuristics keep it readable without any markup: a short line in ALL CAPS or ending in ':' is a
bold heading; lines starting with -, •, *, or · become bullets; blank lines add spacing.
"""
import io
import re

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

_BULLET = re.compile(r"^\s*[-•*·]\s+(.*)$")


def _is_heading(line):
    s = line.strip()
    if not s or len(s) > 60:
        return False
    if s.endswith(":"):
        return True
    letters = [c for c in s if c.isalpha()]
    return bool(letters) and s == s.upper()      # ALL CAPS section header


def build_docx(text, title=""):
    """Return .docx bytes for `text`. Raises ImportError if python-docx is unavailable."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    if title:
        h = doc.add_paragraph()
        run = h.add_run(title)
        run.bold = True
        run.font.size = Pt(15)

    for raw in (text or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            doc.add_paragraph("")
            continue
        m = _BULLET.match(line)
        if m:
            doc.add_paragraph(m.group(1), style="List Bullet")
        elif _is_heading(line):
            p = doc.add_paragraph()
            p.add_run(line.strip()).bold = True
        else:
            doc.add_paragraph(line)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()

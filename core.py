"""
core.py — all the logic for the job-match app, kept free of Streamlit so it can
be tested on its own. app.py imports from here and only handles the UI.
"""

import csv
import os
import re
import json
import math
from io import BytesIO
from collections import Counter

import requests
from bs4 import BeautifulSoup

HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                         "AppleWebKit/537.36 (KHTML, like Gecko) "
                         "Chrome/124.0 Safari/537.36"}

# Default model for AI tailoring; swap to any model string your key can access.
AI_MODEL = "claude-sonnet-4-6"


# ------------------------------------------------------------
# Jobs
# ------------------------------------------------------------
SAMPLE_JOBS = [
    {"found_date": "", "title": "Project Coordinator", "company": "Sample Co",
     "location": "Boston, MA", "url": "https://example.com/jobs/1", "sponsors_h1b": "unknown"},
    {"found_date": "", "title": "Associate Program Manager", "company": "Sample Co",
     "location": "Remote (US)", "url": "https://example.com/jobs/2", "sponsors_h1b": "unknown"},
    {"found_date": "", "title": "Operations Analyst", "company": "Sample Co",
     "location": "New York, NY", "url": "https://example.com/jobs/3", "sponsors_h1b": "unknown"},
]


def load_jobs(path="jobs.csv"):
    """Read the scraper's output. Returns sample rows ONLY when jobs.csv doesn't
    exist yet (so the UI is explorable on first launch). If the file exists but is
    empty, returns [] so the app can explain that filters removed everything."""
    if not os.path.exists(path):
        return list(SAMPLE_JOBS)
    with open(path, newline="", encoding="utf-8") as f:
        return [r for r in csv.DictReader(f) if r.get("url")]


# ------------------------------------------------------------
# Keyword extraction + resume<->JD matching
# ------------------------------------------------------------
STOPWORDS = set("""
a an and or but the of to in for on with at by from as is are was were be been being
this that these those it its their our your you we they he she them his her not no nor
will would can could should may might must do does did done have has had having about
into over under again further then once here there all any both each few more most other
some such only own same so than too very up down out off above below between during who
whom which what when where why how if because while per via etc within across upon also
us am i me my mine ours yours job role position company team work
""".split())

# Words that appear in nearly every JD and carry little signal for matching.
JD_BOILERPLATE = set("""
experience experiences working ability abilities years year skills skill strong excellent
required requirement requirements responsibility responsibilities preferred plus including
include includes related ideal candidate candidates opportunity opportunities environment
help support ensure provide drive build manage manages managing lead leads leading deliver
delivers cross functional looking join growth fast paced paced ll re ve day days week weeks
new like using use used able high level levels world class etc apply applicants applicant
benefits equal employer diversity inclusive
salary base compensation pay hourly bonus equity range eligible eligibility insurance
medical dental vision pto holidays veteran gender race disability accommodation reasonable
location locations remote hybrid onsite office travel sponsorship visa authorization citizen
status eeo position week weekly month monthly annual annually
""".split())

WORD_RE = re.compile(r"[a-zA-Z][a-zA-Z0-9+#./-]{1,}")


def _tokens(text):
    return [w.lower() for w in WORD_RE.findall(text or "")]


_SEGMENT_RE = re.compile(r"[.,;:/()\[\]{}\n\t\u2022|]+")


def extract_keywords(text, top_n=28, max_bigrams=8, extra_skip=None):
    """Pull the most signal-bearing terms from a job description: meaningful
    single words plus two-word phrases. Phrases are formed only WITHIN a clause
    (text is split on punctuation first) so we never glue together words from
    different sentences. `extra_skip` lets the caller drop more words (e.g. the
    company name) so they don't show up as résumé 'gaps'."""
    skip = STOPWORDS | JD_BOILERPLATE | (set(extra_skip) if extra_skip else set())
    uni_counter, bi_counter = Counter(), Counter()
    for segment in _SEGMENT_RE.split(text or ""):
        raw = [t.strip("-.+#/") for t in _tokens(segment)]
        raw = [t for t in raw if t]
        uni_counter.update(t for t in raw if len(t) > 2 and t not in skip)
        for a, b in zip(raw, raw[1:]):          # only truly adjacent words
            if a in skip or b in skip or len(a) <= 2 or len(b) <= 2:
                continue
            bi_counter[f"{a} {b}"] += 1

    bigrams = [k for k, _ in sorted(bi_counter.items(),
                                    key=lambda kv: (kv[1], len(kv[0])), reverse=True)][:max_bigrams]
    bigram_words = {w for bg in bigrams for w in bg.split()}
    unigrams = [k for k, _ in sorted(uni_counter.items(),
                                     key=lambda kv: (kv[1], kv[0]), reverse=True)]

    keywords = list(bigrams)
    for u in unigrams:
        if u in bigram_words:           # already covered by a chosen phrase
            continue
        keywords.append(u)
        if len(keywords) >= top_n:
            break
    return keywords[:top_n]


def match_resume(resume_text, jd_text, top_n=30, extra_skip=None):
    """Return (score 0-100, matched_keywords, missing_keywords)."""
    keywords = extract_keywords(jd_text, top_n=top_n, extra_skip=extra_skip)
    resume_low = (resume_text or "").lower()
    matched, missing = [], []
    for kw in keywords:
        (matched if kw in resume_low else missing).append(kw)
    score = round(100 * len(matched) / len(keywords)) if keywords else 0
    return score, matched, missing


# ------------------------------------------------------------
# Skill-based (semantic-ish) matching — more meaningful than raw word overlap.
# Each canonical skill has aliases; a job's skills = the canonical skills whose
# aliases appear in the JD, and the score = how many of THOSE your resume covers.
# ------------------------------------------------------------
SKILLS = {
    "project management": ("project management", "project manager", "manage projects",
                           "project delivery", "project lifecycle", "project coordination", "pmo", "pmp"),
    "program management": ("program management", "program manager"),
    "agile / scrum": ("agile", "scrum", "kanban", "sprint", "safe"),
    "waterfall": ("waterfall",),
    "stakeholder management": ("stakeholder", "stakeholders", "cross-functional",
                               "cross functional", "relationship management"),
    "risk management": ("risk management", "risk assessment", "risk mitigation", "risk register"),
    "budgeting & cost": ("budget", "budgeting", "cost management", "cost planning",
                         "forecasting", "financial planning", "financial analysis", "p&l"),
    "scheduling": ("scheduling", "timeline", "milestone", "gantt", "ms project",
                   "microsoft project", "primavera", "critical path"),
    "process improvement": ("process improvement", "process optimization", "operational efficiency",
                            "continuous improvement", "lean", "six sigma", "kaizen"),
    "change management": ("change management", "organizational change"),
    "vendor & procurement": ("vendor", "procurement", "supplier", "sourcing", "contract management"),
    "requirements / business analysis": ("requirements gathering", "business analysis",
                                         "business analyst", "brd", "user stories",
                                         "functional specification", "use cases", "gap analysis"),
    "data analysis": ("data analysis", "data analytics", "analytics", "quantitative"),
    "reporting & dashboards": ("reporting", "dashboard", "kpi", "metrics", "visualization"),
    "sql": ("sql", "queries"),
    "python": ("python",),
    "excel": ("excel", "spreadsheet", "pivot table", "vlookup"),
    "power bi": ("power bi", "powerbi"),
    "tableau": ("tableau",),
    "jira": ("jira",),
    "confluence": ("confluence",),
    "asana / trello / wrike": ("asana", "trello", "wrike", "monday.com", "smartsheet", "clickup"),
    "crm / salesforce": ("crm", "salesforce", "hubspot"),
    "erp systems": ("erp", "sap", "netsuite", "workday"),
    "automation": ("automation", "power automate", "workflow automation", "rpa", "zapier"),
    "operations": ("operations", "operational", "logistics", "supply chain"),
    "customer success": ("customer success", "client success", "customer experience",
                         "account management", "onboarding"),
    "quality assurance": ("quality assurance", "quality control", "quality management"),
    "documentation / SOPs": ("documentation", "sop", "standard operating procedure"),
    "leadership": ("leadership", "team lead", "mentoring", "people management", "coaching"),
    "communication": ("communication", "presentation", "presentations"),
    "problem solving": ("problem solving", "problem-solving", "troubleshooting", "analytical thinking"),
    "compliance & audit": ("compliance", "regulatory", "audit", "governance"),
    "product & roadmap": ("product management", "roadmap", "backlog", "prioritization"),
    "go-to-market": ("go-to-market", "gtm", "campaign"),
    "execution & delivery": ("execution", "delivery", "deliverables", "milestones", "on-time", "on time"),
    "prioritization": ("prioritization", "prioritize", "trade-offs", "tradeoffs", "triage"),
    "metrics & kpis": ("kpi", "kpis", "okr", "okrs", "data-driven", "metrics-driven", "key performance"),
    "client-facing": ("client-facing", "client facing", "customer-facing", "liaison", "partner management"),
    "process design": ("process design", "process mapping", "workflow", "operational excellence", "streamline"),
}

_SKILL_RES = {c: re.compile(r"\b(?:%s)\b" % "|".join(re.escape(a) for a in aliases), re.I)
              for c, aliases in SKILLS.items()}


def skills_in(text):
    """Canonical skills whose aliases appear in `text`."""
    text = text or ""
    return {c for c, rx in _SKILL_RES.items() if rx.search(text)}


# Broad skills that show up in almost every JD — down-weighted so matching them
# doesn't inflate the score the way specific skills (SQL, risk management, Jira) do.
SKILL_WEIGHTS = {
    "operations": 0.4, "communication": 0.4, "leadership": 0.4, "problem solving": 0.5,
    "reporting & dashboards": 0.6, "documentation / SOPs": 0.5, "customer success": 0.6,
    "quality assurance": 0.6, "go-to-market": 0.6, "product & roadmap": 0.7,
}

_IDF_PATH = "idf.json"


def _candidate_terms(text):
    """All meaningful unigrams + adjacent bigrams in a text (for IDF / matching)."""
    skip = STOPWORDS | JD_BOILERPLATE
    terms = set()
    for segment in _SEGMENT_RE.split(text or ""):
        raw = [t.strip("-.+#/") for t in _tokens(segment)]
        raw = [t for t in raw if t]
        for t in raw:
            if len(t) > 2 and t not in skip:
                terms.add(t)
        for a, b in zip(raw, raw[1:]):
            if a not in skip and b not in skip and len(a) > 2 and len(b) > 2:
                terms.add("%s %s" % (a, b))
    return terms


def build_idf(texts):
    """Inverse document frequency over a corpus of JDs: common terms get a low
    weight, rare/important terms a high one."""
    texts = list(texts)
    n = len(texts) or 1
    df = Counter()
    for t in texts:
        for term in _candidate_terms(t):
            df[term] += 1
    return {term: round(math.log((n + 1) / (c + 1)) + 1.0, 4) for term, c in df.items()}


def save_idf(idf, path=_IDF_PATH):
    try:
        json.dump(idf, open(path, "w", encoding="utf-8"))
    except Exception:
        pass


def load_idf(path=_IDF_PATH):
    if os.path.exists(path):
        try:
            return json.load(open(path, encoding="utf-8"))
        except Exception:
            return None
    return None


_REQ_HEADERS = ("minimum qualifications", "basic qualifications", "preferred qualifications",
                "qualifications", "requirements", "what you'll need", "what we're looking for",
                "who you are", "what you'll do", "responsibilities", "skills you")


def _requirements_text(jd_text):
    """The part of a JD from its first 'requirements/qualifications/responsibilities'
    header onward (where the real must-haves live). '' if none found."""
    low = (jd_text or "").lower()
    hits = [low.find(h) for h in _REQ_HEADERS if low.find(h) != -1]
    return jd_text[min(hits):] if hits else ""


# Hard skills, tools, methods, and certs an ATS literally scans for — weighted highest.
ATS_KEYWORDS = {
    # tools
    "jira", "confluence", "asana", "trello", "smartsheet", "monday.com", "wrike", "clickup",
    "ms project", "microsoft project", "primavera", "sharepoint", "excel", "google sheets",
    "powerpoint", "visio", "lucidchart", "miro", "notion", "sql", "tableau", "power bi",
    "looker", "salesforce", "sap", "oracle", "netsuite", "workday", "servicenow", "python", "git",
    # methods / frameworks
    "agile", "scrum", "kanban", "safe", "lean", "six sigma", "lean six sigma", "waterfall",
    "sdlc", "devops", "kaizen", "pmbok", "prince2", "itil", "okr", "okrs", "kpi", "kpis",
    "gantt", "sprint", "backlog", "retrospective", "scrum master", "product owner",
    # certifications
    "pmp", "capm", "csm", "psm", "cspo", "cbap", "green belt", "black belt",
    # PM / analyst / ops domain
    "project management", "program management", "project manager", "program manager",
    "project coordinator", "stakeholder management", "stakeholder", "risk management",
    "change management", "budget", "budgeting", "cost management", "resource allocation",
    "scope management", "requirements gathering", "business requirements", "user stories",
    "process improvement", "process mapping", "gap analysis", "data analysis", "reporting",
    "dashboards", "forecasting", "vendor management", "procurement", "milestones",
    "deliverables", "cross-functional", "roadmap", "status reporting", "project plan",
    "business analysis", "operations", "implementation", "onboarding", "sla", "metrics",
}


def skill_match(resume_text, jd_text, idf=None):
    """ATS-style match: score = weighted % of the JD's important keywords present in the
    resume — hard skills / tools / certs and requirement-section terms weighted highest,
    exactly how a keyword-screening ATS works. Returns (score, keywords_have, keywords_to_add)."""
    resume_low = (resume_text or "").lower()
    jd_low = (jd_text or "").lower()
    req_low = _requirements_text(jd_text).lower()

    # The JD's important keywords: its salient terms + any hard ATS keywords it names.
    jd_terms = set(extract_keywords(jd_text, top_n=30))
    jd_terms |= {kw for kw in ATS_KEYWORDS if kw in jd_low}
    if not jd_terms:
        return 0, [], []

    default_w = max(idf.values()) if idf else 1.0

    def wt(t):
        w = idf.get(t, default_w) if idf else 1.0
        if t in ATS_KEYWORDS:        # hard skill / tool / cert — what an ATS weights most
            w *= 2.5
        if t in req_low:             # stated in the requirements/qualifications section
            w *= 1.6
        return w

    have = sorted((t for t in jd_terms if t in resume_low), key=lambda t: -wt(t))
    missing = sorted((t for t in jd_terms if t not in resume_low), key=lambda t: -wt(t))
    score = round(100.0 * sum(wt(t) for t in have) / (sum(wt(t) for t in jd_terms) or 1.0))
    return score, have, missing


# ------------------------------------------------------------
# Sponsorship signal — the single biggest time-saver for an international student.
# A company can be a known H-1B sponsor yet post a role that explicitly WON'T work
# for a visa candidate (no sponsorship, or it needs citizenship / a clearance / a
# green card). We read that straight from the JD so those roles can be flagged/hidden.
# ------------------------------------------------------------
_SPONSOR_BLOCK = [
    ("no_sponsor", "JD says no visa sponsorship", re.compile(
        r"(?:will|are|is|can|do(?:es)?)?\s*(?:not|n't|unable|never)\b[^.]{0,40}\bsponsor"
        r"|\bno\b[^.]{0,15}\bsponsorship"
        r"|\bwithout[^.]{0,30}\bsponsorship"
        r"|\bsponsorship[^.]{0,20}\bnot\b[^.]{0,20}(?:available|offered|provided|considered)"
        r"|\bnot[^.]{0,15}(?:offer|provide|consider)[^.]{0,15}sponsorship"
        r"|\bdo(?:es)? not (?:require|need)[^.]{0,20}sponsorship"
        r"|authoriz(?:ed|ation) to work[^.]{0,70}without[^.]{0,25}sponsor", re.I)),
    ("citizen", "JD requires U.S. citizenship", re.compile(
        r"\bmust be (?:a |an )?(?:u\.?s\.?\s*)?citizen"
        r"|\b(?:u\.?s\.?\s*)?citizenship\b[^.]{0,20}\b(?:is required|required|requirement|mandatory|only)"
        r"|\b(?:require[sd]?|requiring)\b.{0,25}?\bcitizenship", re.I)),
    ("clearance", "JD requires a security clearance", re.compile(
        r"\b(?:security|government)\b[^.]{0,15}clearance"
        r"|\bactive[^.]{0,20}clearance"
        r"|\bts/sci\b|\btop secret\b|\bsecret clearance\b|\bpublic trust\b"
        r"|\bclearance (?:is )?(?:required|eligible|active)", re.I)),
    ("greencard", "JD requires a green card / permanent residency", re.compile(
        r"\b(?:green card|permanent residen(?:t|cy|ce))\b[^.]{0,25}(?:require|must|only|holder)"
        r"|must be (?:a )?(?:green card holder|permanent resident)", re.I)),
]
_SPONSOR_OPEN = re.compile(
    r"(?:visa|h-?1b|employment|work)?\s*sponsorship (?:is |may be |are |can be )?"
    r"(?:available|offered|provided|possible|considered|supported)"
    r"|(?:will|can|do|happy to|open to|able to|willing to|we)\s+(?:gladly |certainly )?sponsor\b"
    r"|(?:offer|provide|support)[^.]{0,20}(?:visa |h-?1b )?sponsorship"
    r"|\bh-?1b[^.]{0,15}sponsorship"
    r"|\bsponsor[^.]{0,15}(?:visa|h-?1b|work authorization)", re.I)


def sponsorship_from_jd(jd_text):
    """Read a JD for an explicit sponsorship signal. Returns (verdict, reason):
      'blocked' = the JD says it won't work for a visa candidate (no sponsorship, or it
                  requires U.S. citizenship / a security clearance / a green card)
      'open'    = the JD explicitly offers visa sponsorship
      ''        = no clear signal (most postings)
    Checks the 'blocked' phrasings first since those are the ones that waste your time."""
    jd = jd_text or ""
    if not jd:
        return "", ""
    for _cat, msg, rx in _SPONSOR_BLOCK:
        if rx.search(jd):
            return "blocked", msg
    if _SPONSOR_OPEN.search(jd):
        return "open", "JD offers visa sponsorship"
    return "", ""


# ------------------------------------------------------------
# H-1B cap-exempt employers (universities, nonprofit hospitals, research institutes).
# Cap-exempt = NO H-1B lottery — a major edge for an international student, so we badge it.
# ------------------------------------------------------------
_CAP_EXEMPT_RE = re.compile(
    r"\b(?:universit(?:y|ies)|college|polytechnic|institute of technology"
    r"|school of (?:medicine|public health|nursing|engineering|law)"
    r"|cancer (?:institute|center|centre)|medical (?:center|centre|college|school)"
    r"|health system|hospital|children'?s hospital|clinic"
    r"|national lab(?:oratory)?|research institute)\b", re.I)
_CAP_EXEMPT_NAMES = ("mayo clinic", "cleveland clinic", "kaiser permanente", "dana-farber",
                     "memorial sloan", "md anderson", "mass general", "brigham and women",
                     "national institutes of health")


def is_cap_exempt(company):
    """Heuristic: True if the employer is LIKELY H-1B cap-exempt (universities, nonprofit
    hospitals, research institutes) → no H-1B lottery. A hint to verify, not a guarantee."""
    c = (company or "").lower()
    if not c:
        return False
    if any(n in c for n in _CAP_EXEMPT_NAMES):
        return True
    return bool(_CAP_EXEMPT_RE.search(c))


# ------------------------------------------------------------
# Sponsor STRENGTH — turn the yes/no flag into a confidence tier using DOL filing
# VOLUME (a company that files thousands of H-1Bs is a far safer bet than one with two).
# Needs an optional sponsor_counts.json {normalized_name: count}; degrades to '' without it.
# ------------------------------------------------------------
def load_sponsor_counts(path="sponsor_counts.json"):
    """Optional {normalized_company: H1B_filing_count} built from DOL LCA data.
    Returns {} when the file is absent (strength just isn't shown)."""
    if not os.path.exists(path):
        return {}
    try:
        return json.load(open(path, encoding="utf-8")) or {}
    except Exception:
        return {}


def sponsor_strength(company, counts):
    """Tier a sponsor by filing VOLUME. Returns ('high'|'medium'|'low'|'', count).
    ('', 0) when there's no number for the company. Counts come from load_sponsor_counts()."""
    if not counts or not company:
        return "", 0
    try:
        import scraper                      # lazy: scraper imports core (avoid circular at load)
        key = scraper._norm_name(company)
    except Exception:
        key = re.sub(r"[^a-z0-9 ]+", " ", company.lower()).strip()
    try:
        n = int(counts.get(key) or counts.get(company.lower()) or 0)
    except Exception:
        n = 0
    if n >= 1000:
        return "high", n
    if n >= 100:
        return "medium", n
    if n >= 1:
        return "low", n
    return "", 0


# ------------------------------------------------------------
# Experience requirement parsing (to keep only entry-level roles)
# ------------------------------------------------------------
_EXP_YEARS_RE = re.compile(r"(\d{1,2})\s*\+\s*years?", re.I)   # 'N+ years' (e.g. '5+ years')


def required_years(text):
    """Best-effort: the highest 'N+ years ... experience' requirement mentioned in the
    text (0 if none). Used to drop roles that demand more experience than you have."""
    if not text:
        return 0
    yrs = [int(m.group(1)) for m in _EXP_YEARS_RE.finditer(text)]
    return max(yrs) if yrs else 0


# ------------------------------------------------------------
# Fetch a job description page (best-effort; paste fallback in the UI)
# ------------------------------------------------------------
def fetch_jd(url, limit=8000):
    try:
        r = requests.get(url, headers=HEADERS, timeout=20)
        r.raise_for_status()
        soup = BeautifulSoup(r.text, "lxml")
        for tag in soup(["script", "style", "nav", "header", "footer", "form"]):
            tag.decompose()
        text = soup.get_text(separator=" ", strip=True)
        return re.sub(r"\s{2,}", " ", text)[:limit]
    except Exception:
        return ""


# ------------------------------------------------------------
# Export the (edited) resume to .docx
# ------------------------------------------------------------
def resume_to_docx_bytes(text):
    """Turn plain-text resume into a simple .docx. ALL-CAPS short lines become
    bold headings; lines starting with - or • become bullets."""
    from docx import Document  # python-docx
    doc = Document()
    for raw in (text or "").split("\n"):
        s = raw.strip()
        if not s:
            doc.add_paragraph("")
        elif s.isupper() and len(s) <= 40:
            p = doc.add_paragraph()
            p.add_run(s).bold = True
        elif s[0] in "-•*":
            doc.add_paragraph(s.lstrip("-•* ").strip(), style="List Bullet")
        else:
            doc.add_paragraph(s)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------
# Optional: tailor the resume to a JD with Claude
# (needs `pip install anthropic` and ANTHROPIC_API_KEY set)
# ------------------------------------------------------------
def ai_available():
    return bool(os.environ.get("ANTHROPIC_API_KEY"))


def tailor_with_ai(resume_text, jd_text, model=AI_MODEL, api_key=None):
    from anthropic import Anthropic
    client = Anthropic(api_key=api_key) if api_key else Anthropic()  # else reads the env var
    prompt = (
        "You are helping a job seeker tailor their resume to one specific job "
        "description. Rewrite the resume so it surfaces the experience and skills "
        "most relevant to the job and mirrors the job's terminology WHERE THE "
        "CANDIDATE GENUINELY HAS THAT EXPERIENCE.\n\n"
        "Hard rules:\n"
        "- Do NOT invent or exaggerate experience, employers, titles, dates, or metrics.\n"
        "- Only reorder, reword, and re-emphasize what is already in the resume.\n"
        "- Keep it concise, truthful, and ATS-friendly.\n"
        "- Return ONLY the revised resume text, no commentary.\n\n"
        f"=== JOB DESCRIPTION ===\n{jd_text}\n\n"
        f"=== CURRENT RESUME ===\n{resume_text}\n\n"
        "=== REVISED RESUME ==="
    )
    msg = client.messages.create(
        model=model,
        max_tokens=2000,
        messages=[{"role": "user", "content": prompt}],
    )
    return "".join(b.text for b in msg.content if getattr(b, "type", "") == "text").strip()
